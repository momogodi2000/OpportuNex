"""CV text extraction and deterministic item extraction (module M03, EF-03.01..04).

Text extraction runs in a separate, short-lived process (:func:`extract_text_isolated`) so a
malformed or hostile file cannot crash or hang the engine (EF-03.02). Only PDF, DOCX and
plain-text files are accepted; no macro or active content is ever executed — PDF text is
read with pdfplumber/pypdf and DOCX through python-docx, which only read XML parts.
"""

from __future__ import annotations

import io
import multiprocessing as mp
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from core.errors import ValidationFailed
from core.text import fold

MAX_BYTES = 10 * 1024 * 1024
PARSE_TIMEOUT_SECONDS = 30
MAX_TEXT_CHARS = 200_000

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TXT = "text/plain"


def detect_mime(name: str, data: bytes) -> str:
    """Validate type by extension *and* content signature."""
    if not data:
        raise ValidationFailed("empty file", code="file_empty")
    if len(data) > MAX_BYTES:
        raise ValidationFailed("file larger than 10 MB", code="file_too_large")
    lower = name.lower()
    if lower.endswith(".pdf"):
        if not data.startswith(b"%PDF-"):
            raise ValidationFailed("not a PDF file", code="file_type_mismatch")
        return PDF
    if lower.endswith(".docx"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
        except zipfile.BadZipFile as exc:
            raise ValidationFailed("not a DOCX file", code="file_type_mismatch") from exc
        if "word/document.xml" not in names:
            raise ValidationFailed("not a DOCX file", code="file_type_mismatch")
        if any(n.lower().endswith("vbaproject.bin") for n in names):
            raise ValidationFailed("macro-enabled documents are refused", code="file_active_content")
        return DOCX
    if lower.endswith(".txt"):
        try:
            data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                data.decode("cp1252")
            except UnicodeDecodeError as exc:
                raise ValidationFailed("text file encoding not supported", code="file_type_mismatch") from exc
        return TXT
    raise ValidationFailed("only PDF, DOCX and TXT are accepted", code="file_type_unsupported")


def _extract_text(mime: str, data: bytes) -> str:
    if mime == PDF:
        import pdfplumber

        parts: list[str] = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:30]:
                parts.append(page.extract_text() or "")
        text = "\n".join(parts)
        if not text.strip():
            from pypdf import PdfReader

            text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages[:30])
        return text
    if mime == DOCX:
        import docx

        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252")


def _worker(mime: str, data: bytes, out: Any) -> None:
    try:
        out.put(("ok", _extract_text(mime, data)[:MAX_TEXT_CHARS]))
    except Exception as exc:  # noqa: BLE001 - reported to the parent, never raised there
        out.put(("error", type(exc).__name__))


def extract_text_isolated(mime: str, data: bytes, timeout: float = PARSE_TIMEOUT_SECONDS) -> str:
    ctx = mp.get_context("spawn")
    queue = ctx.Queue()
    proc = ctx.Process(target=_worker, args=(mime, data, queue), daemon=True)
    proc.start()
    try:
        status, payload = queue.get(timeout=timeout)
    except Exception as exc:
        proc.kill()
        raise ValidationFailed("document could not be read in time", code="file_parse_timeout") from exc
    finally:
        proc.join(timeout=2)
        if proc.is_alive():
            proc.kill()
    if status != "ok":
        raise ValidationFailed("document is malformed", code="file_malformed")
    return str(payload)


# ---------------------------------------------------------------- extraction
SECTION_PATTERNS = {
    "education": r"(formation|formations|education|etudes|parcours academique|diplomes?|academic background)",
    "experience": r"(experiences?( professionnelles?)?|work experience|professional experience|emplois?|employment"
    r"|parcours professionnel)",
    "skills": r"(competences|skills|savoir[- ]faire|competences techniques|technical skills|outils|tools)",
    "languages": r"(langues|languages|langues parlees)",
    "certifications": r"(certifications?|certificats?|licences? et certifications)",
}
_SECTION_RE = {k: re.compile(rf"^\s*{v}\s*:?\s*$") for k, v in SECTION_PATTERNS.items()}

_YEAR = r"(19|20)\d{2}"
_MONTHS = (
    r"(janv|f[eé]vr|mars|avr|mai|juin|juil|ao[uû]t|sept|oct|nov|d[eé]c|jan|feb|mar|apr|may|jun|jul|"
    r"aug|sep|[01]?\d)[a-zéû.]*"
)
_DATE = rf"(?:{_MONTHS}[\s/.-]*)?{_YEAR}"
_RANGE_RE = re.compile(
    rf"(?P<start>{_DATE})\s*(?:-|–|—|a|à|to|au)\s*(?P<end>{_DATE}|pr[eé]sent|aujourd['’]?hui|"
    r"actuel(?:lement)?|now|current|en cours)",
    re.IGNORECASE,
)
_SINGLE_YEAR_RE = re.compile(rf"\b{_YEAR}\b")
_URL_RE = re.compile(r"https?://[^\s)>\]]+", re.IGNORECASE)
_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")

LANGUAGE_NAMES = {
    "francais": "fr",
    "french": "fr",
    "anglais": "en",
    "english": "en",
    "espagnol": "es",
    "spanish": "es",
    "portugais": "pt",
    "portuguese": "pt",
    "allemand": "de",
    "german": "de",
    "arabe": "ar",
    "arabic": "ar",
    "chinois": "zh",
    "chinese": "zh",
    "mandarin": "zh",
    "swahili": "sw",
}
LEVEL_WORDS = [
    (r"\b(c2|bilingue|bilingual|courant|fluent)\b", "C2"),
    (r"\bc1\b|\bavance|\badvanced\b", "C1"),
    (r"\bb2\b|\bintermediaire avance|upper[- ]intermediate", "B2"),
    (r"\bb1\b|\bintermediaire|\bintermediate", "B1"),
    (r"\ba2\b|\bscolaire|\belementaire|\belementary", "A2"),
    (r"\ba1\b|\bdebutant|\bbeginner|\bnotions?\b", "A1"),
    (r"\b(maternelle|native|langue maternelle|mother tongue)\b", "native"),
]
DEGREE_WORDS = re.compile(
    r"\b(master|licence|bachelor|doctorat|phd|bts|dut|hnd|baccalaur[ée]at|bac\b|mba|"
    r"ingenieur|ing[ée]nieur|dipl[ôo]me|degree|msc|bsc|ma\b|ba\b)",
    re.IGNORECASE,
)


@dataclass
class ExtractedItem:
    item_type: str
    payload: dict[str, Any]
    confidence: float


@dataclass
class CvExtraction:
    items: list[ExtractedItem] = field(default_factory=list)


def _norm_date(token: str | None) -> str | None:
    if not token:
        return None
    t = fold(token)
    if re.search(r"present|aujourd|actuel|now|current|en cours", t):
        return None
    year = re.search(_YEAR, t)
    if not year:
        return None
    months = {
        "janv": 1,
        "jan": 1,
        "fevr": 2,
        "feb": 2,
        "mars": 3,
        "mar": 3,
        "avr": 4,
        "apr": 4,
        "mai": 5,
        "may": 5,
        "juin": 6,
        "jun": 6,
        "juil": 7,
        "jul": 7,
        "aout": 8,
        "aug": 8,
        "sept": 9,
        "sep": 9,
        "oct": 10,
        "nov": 11,
        "dec": 12,
    }
    month = None
    m = re.match(r"([a-z]+)", t)
    if m:
        month = next((v for k, v in months.items() if m.group(1).startswith(k)), None)
    else:
        num = re.match(r"(\d{1,2})[\s/.-]", t)
        if num and 1 <= int(num.group(1)) <= 12:
            month = int(num.group(1))
    return f"{year.group(0)}-{month:02d}" if month else year.group(0)


def _split_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in lines:
        f = fold(line)
        found = next((k for k, rx in _SECTION_RE.items() if rx.match(f)), None)
        if found:
            current = found
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def _is_date_only(line: str) -> bool:
    return bool(_RANGE_RE.search(line)) and len(_RANGE_RE.sub("", line).strip(" -–—|,:()")) < 3


def _blocks(lines: list[str]) -> list[list[str]]:
    """Group a section into entries anchored on date ranges.

    Handles "Title – Org (2019-2021)", "Title – Org" followed by a date line, and a date
    line followed by the title.
    """
    entries: list[list[str]] = []
    pending: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if not _RANGE_RE.search(line):
            pending.append(line)
            continue
        if _is_date_only(line):
            if pending:
                title = pending.pop()
                if entries:
                    entries[-1].extend(pending)
                entries.append([title, line])
            else:
                entries.append([line])
        else:
            if entries:
                entries[-1].extend(pending)
            entries.append([line])
        pending = []
    if entries:
        entries[-1].extend(pending)
    elif pending:
        entries.append(pending)
    return entries


def _title_and_rest(block: list[str]) -> tuple[str, list[str]]:
    content = [ln for ln in block if not _is_date_only(ln)]
    if not content:
        return "", []
    return _strip_dates(content[0]), content[1:]


def _dates(text: str) -> tuple[str | None, str | None, bool]:
    m = _RANGE_RE.search(text)
    if m:
        end_raw = m.group("end")
        ongoing = _norm_date(end_raw) is None
        return _norm_date(m.group("start")), _norm_date(end_raw), ongoing
    years = [y.group(0) for y in _SINGLE_YEAR_RE.finditer(text)]
    return (None, years[-1], False) if years else (None, None, False)


def _strip_dates(text: str) -> str:
    text = _RANGE_RE.sub("", text)
    return re.sub(r"\s{2,}", " ", text).strip(" -–—|,:")


def extract_items(text: str, known_skills: dict[str, str]) -> CvExtraction:
    """``known_skills`` maps folded alias → canonical skill name."""
    lines = [ln.rstrip() for ln in text.replace("\r", "").split("\n")]
    sections = _split_sections(lines)
    out = CvExtraction()

    header = [ln.strip() for ln in sections.get("header", []) if ln.strip()]
    if header:
        first = header[0]
        if 2 <= len(first.split()) <= 5 and not _EMAIL_RE.search(first) and not any(ch.isdigit() for ch in first):
            out.items.append(ExtractedItem("identity", {"full_name": first.title() if first.isupper() else first}, 0.6))

    for url in dict.fromkeys(_URL_RE.findall(text)):
        kind = (
            "linkedin"
            if "linkedin." in url
            else "github"
            if "github." in url
            else "gitlab"
            if "gitlab." in url
            else "orcid"
            if "orcid.org" in url
            else "website"
        )
        out.items.append(ExtractedItem("link", {"kind": kind, "url": url.rstrip(".,;")}, 0.9))

    for block in _blocks(sections.get("education", [])):
        joined = " ".join(block)
        start, end, ongoing = _dates(joined)
        title_line, rest = _title_and_rest(block)
        parts = [p.strip() for p in re.split(r"\s[-–—|,]\s|\s*\|\s*", title_line) if p.strip()]
        degree = next((p for p in parts if DEGREE_WORDS.search(p)), parts[0] if parts else title_line)
        institution = next((p for p in parts if p != degree), rest[0] if rest else "")
        level_match = DEGREE_WORDS.search(degree)
        out.items.append(
            ExtractedItem(
                "education",
                {
                    "level": level_match.group(0).title() if level_match else degree[:60],
                    "field_of_study": degree[:160],
                    "institution": _strip_dates(institution)[:200] or "—",
                    "start_date": start,
                    "end_date": None if ongoing else end,
                    "is_ongoing": ongoing,
                },
                0.7 if level_match and institution else 0.45,
            )
        )

    for block in _blocks(sections.get("experience", [])):
        joined = " ".join(block)
        start, end, ongoing = _dates(joined)
        title_line, rest = _title_and_rest(block)
        parts = [p.strip() for p in re.split(r"\s[-–—|,@]\s|\s+chez\s+|\s+at\s+|\s*\|\s*", title_line) if p.strip()]
        job = parts[0] if parts else title_line
        employer = parts[1] if len(parts) > 1 else (rest[0] if rest else "")
        description = " ".join(rest if len(parts) > 1 else rest[1:])
        if start is None:
            continue  # an experience without a start date cannot be stored (schema NOT NULL)
        out.items.append(
            ExtractedItem(
                "experience",
                {
                    "job_title": job[:160],
                    "employer": _strip_dates(employer)[:200] or "—",
                    "start_date": start,
                    "end_date": None if ongoing else end,
                    "is_current": ongoing,
                    "description": description[:4000] or None,
                },
                0.7 if employer else 0.45,
            )
        )

    skill_text = " ".join(sections.get("skills", [])) or text
    found_skills: dict[str, float] = {}
    folded_all = f" {fold(skill_text)} "
    for alias, name in known_skills.items():
        if len(alias) >= 2 and re.search(rf"(?<![\w]){re.escape(alias)}(?![\w])", folded_all):
            found_skills[name] = max(found_skills.get(name, 0), 0.8 if sections.get("skills") else 0.55)
    for raw in re.split(r"[,;•·\n]", "\n".join(sections.get("skills", []))):
        token = raw.strip(" -*\t")
        if 2 <= len(token) <= 40 and len(token.split()) <= 4 and fold(token) not in known_skills:
            found_skills.setdefault(token, 0.5)
    for name, conf in found_skills.items():
        out.items.append(ExtractedItem("skill", {"name": name}, conf))

    lang_lines = sections.get("languages", []) or [ln for ln in lines if any(n in fold(ln) for n in LANGUAGE_NAMES)]
    seen_lang: set[str] = set()
    for line in lang_lines:
        f = fold(line)
        for word, code in LANGUAGE_NAMES.items():
            if re.search(rf"\b{word}\b", f) and code not in seen_lang:
                seen_lang.add(code)
                level = next((lvl for rx, lvl in LEVEL_WORDS if re.search(rx, f)), None)
                out.items.append(
                    ExtractedItem(
                        "language", {"language_code": code, "cefr_level": level or "B1"}, 0.75 if level else 0.4
                    )
                )

    for line in sections.get("certifications", []):
        if line.strip():
            _s, end, _o = _dates(line)
            out.items.append(
                ExtractedItem("certification", {"name": _strip_dates(line)[:200], "obtained_on": end}, 0.6)
            )
    return out
