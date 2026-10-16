"""Exports (module M16): CSV, XLSX, PDF report, versioned JSON, full personal-data archive."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import date
from typing import Any

from sqlalchemy import insert, select

from core.errors import ValidationFailed
from core.ids import uuid7
from core.timeutil import now_iso
from core.version import ENGINE_VERSION
from services.state import ActiveProfile

JSON_SCHEMA_VERSION = "1.0"
COLUMNS = [
    ("title", "Titre"),
    ("organization_name", "Organisme"),
    ("category", "Catégorie"),
    ("country_code", "Pays"),
    ("deadline_on", "Date limite"),
    ("amount", "Montant"),
    ("score", "Pertinence"),
    ("canonical_url", "Lien officiel"),
    ("verification_status", "Statut"),
]
PERSONAL_COLUMNS = [("status", "Statut candidature"), ("priority", "Priorité"), ("notes", "Notes")]
PERSONAL_TABLES = (
    "users",
    "user_profiles",
    "profile_preferred_locations",
    "education_records",
    "certifications",
    "work_experiences",
    "user_skills",
    "user_languages",
    "profile_links",
    "user_objectives",
    "user_custom_objectives",
    "project_profiles",
    "organizations",
    "organization_members",
    "research_workspaces",
    "search_queries",
    "search_runs",
    "workspace_opportunities",
    "workspace_notes",
    "matching_weights",
    "match_feedback",
    "saved_opportunities",
    "tags",
    "saved_opportunity_tags",
    "applications",
    "application_status_history",
    "application_events",
    "application_notes",
    "checklist_items",
    "application_documents",
    "generated_documents",
    "generated_document_versions",
    "reminders",
    "monitors",
    "monitor_sources",
    "notification_records",
    "notification_preferences",
    "consents",
    "data_disclosures",
    "ai_providers",
    "ai_task_routing",
    "export_records",
)
SECRET_COLUMNS = {"users": {"password_hash", "recovery_key_hash", "wrapped_db_key"}}


def _money(row: dict[str, Any]) -> str:
    if not row.get("amount_currency"):
        return ""
    factor = 1 if row["amount_currency"] in ("XAF", "XOF") else 100
    lo = (row.get("amount_min_minor") or 0) / factor
    hi = (row.get("amount_max_minor") or 0) / factor
    return f"{lo:,.0f}{'–' + format(hi, ',.0f') if hi else ''} {row['amount_currency']}".replace(",", " ")


class ExportService:
    def __init__(self, active: ActiveProfile) -> None:
        self.active = active
        self.T = active.tables
        self.uid = active.user_id

    # ---------------------------------------------------------- selection
    def _rows(
        self, scope: str, ids: list[str] | None, run_id: str | None, workspace_id: str | None, include_personal: bool
    ) -> list[dict[str, Any]]:
        T = self.T
        with self.active.db.read() as conn:
            if scope == "selection":
                if not ids:
                    raise ValidationFailed("nothing selected", code="export_empty_selection")
                opp_ids = list(ids)
            elif scope == "filtered_results":
                if not run_id:
                    raise ValidationFailed("run_id required", code="export_run_required")
                owner = conn.execute(select(T.search_runs.c.user_id).where(T.search_runs.c.id == run_id)).scalar()
                if owner != self.uid:
                    raise ValidationFailed("unknown run", code="run_not_found")
                opp_ids = list(
                    conn.execute(
                        select(T.search_results.c.opportunity_id)
                        .where(T.search_results.c.search_run_id == run_id)
                        .order_by(T.search_results.c.rank)
                    ).scalars()
                )
            elif scope == "workspace":
                opp_ids = list(
                    conn.execute(
                        select(T.workspace_opportunities.c.opportunity_id).where(
                            T.workspace_opportunities.c.workspace_id == workspace_id
                        )
                    ).scalars()
                )
                opp_ids += list(
                    conn.execute(
                        select(T.saved_opportunities.c.opportunity_id).where(
                            T.saved_opportunities.c.workspace_id == workspace_id,
                            T.saved_opportunities.c.user_id == self.uid,
                        )
                    ).scalars()
                )
            elif scope in ("saved", "applications"):
                opp_ids = list(
                    conn.execute(
                        select(T.saved_opportunities.c.opportunity_id).where(
                            T.saved_opportunities.c.user_id == self.uid
                        )
                    ).scalars()
                )
            else:
                raise ValidationFailed("unknown scope", code="export_scope_invalid")
            opp_ids = list(dict.fromkeys(opp_ids))
            rows = {
                r["id"]: dict(r)
                for r in conn.execute(select(T.opportunities).where(T.opportunities.c.id.in_(opp_ids))).mappings()
            }
            cats = dict(conn.execute(select(T.opportunity_categories.c.id, T.opportunity_categories.c.code)).all())
            scores = dict(
                conn.execute(
                    select(T.match_scores.c.opportunity_id, T.match_scores.c.score).where(
                        T.match_scores.c.user_id == self.uid
                    )
                ).all()
            )
            personal: dict[str, dict[str, Any]] = {}
            if include_personal or scope == "applications":
                for r in conn.execute(
                    select(
                        T.saved_opportunities.c.opportunity_id,
                        T.saved_opportunities.c.priority,
                        T.applications.c.status,
                        T.saved_opportunities.c.id,
                    )
                    .join(T.applications, T.applications.c.saved_id == T.saved_opportunities.c.id)
                    .where(T.saved_opportunities.c.user_id == self.uid)
                ).all():
                    notes = (
                        conn.execute(select(T.application_notes.c.body).where(T.application_notes.c.saved_id == r[3]))
                        .scalars()
                        .all()
                    )
                    personal[r[0]] = {"priority": r[1], "status": r[2], "notes": " | ".join(notes)}
        out = []
        for oid in opp_ids:
            if oid not in rows:
                continue
            row = rows[oid]
            item = {
                "id": oid,
                "title": row["title"],
                "organization_name": row["organization_name"] or "",
                "category": cats.get(row["category_id"], ""),
                "country_code": row["country_code"] or "",
                "deadline_on": row["deadline_on"] or "",
                "amount": _money(row),
                "score": scores.get(oid, ""),
                "canonical_url": row["canonical_url"],
                "verification_status": row["verification_status"],
                "summary": row["summary"] or "",
                "eligibility_raw": row["eligibility_raw"] or "",
                "deadline_raw": row["deadline_raw"] or "",
            }
            if oid in personal:
                item |= personal[oid]
            out.append(item)
        return out

    def _record(self, fmt: str, scope: str, filters: dict[str, Any], count: int, personal: bool) -> None:
        with self.active.db.tx() as conn:
            conn.execute(
                insert(self.T.export_records).values(
                    id=uuid7(),
                    user_id=self.uid,
                    format=fmt,
                    scope=scope,
                    filters=json.dumps(filters),
                    item_count=count,
                    includes_personal_data=int(personal),
                )
            )

    def export(
        self,
        fmt: str,
        scope: str,
        *,
        ids: list[str] | None = None,
        run_id: str | None = None,
        workspace_id: str | None = None,
        include_personal: bool = False,
        lang: str = "fr",
    ) -> tuple[str, str, bytes]:
        rows = self._rows(scope, ids, run_id, workspace_id, include_personal)
        personal = include_personal or scope == "applications"
        cols = COLUMNS + (PERSONAL_COLUMNS if personal else [])
        stamp = date.today().isoformat()
        self._record(
            fmt, scope, {"run_id": run_id, "workspace_id": workspace_id, "count": len(rows)}, len(rows), personal
        )
        name = f"opportunex-{scope}-{stamp}"
        if fmt == "csv":
            buf = io.StringIO()
            w = csv.writer(buf, delimiter=";")
            w.writerow([label for _k, label in cols])
            for r in rows:
                w.writerow([r.get(k, "") for k, _ in cols])
            return f"{name}.csv", "text/csv; charset=utf-8", ("﻿" + buf.getvalue()).encode("utf-8")
        if fmt == "xlsx":
            return (
                f"{name}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                self._xlsx(rows, cols),
            )
        if fmt == "json":
            doc = {
                "schema": "https://opportunex.app/schemas/export",
                "schema_version": JSON_SCHEMA_VERSION,
                "generator": f"Opportunex {ENGINE_VERSION}",
                "exported_at": now_iso(),
                "scope": scope,
                "includes_personal_data": personal,
                "items": rows,
            }
            return f"{name}.json", "application/json", json.dumps(doc, ensure_ascii=False, indent=2).encode("utf-8")
        if fmt == "pdf":
            return f"{name}.pdf", "application/pdf", self._pdf(rows, scope, lang)
        raise ValidationFailed("unknown format", code="format_invalid")

    @staticmethod
    def _xlsx(rows: list[dict[str, Any]], cols: list[tuple[str, str]]) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Opportunités"
        ws.append([label for _k, label in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E5BC6")
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append([r.get(k, "") for k, _ in cols])
            link = ws.cell(row=ws.max_row, column=[k for k, _ in cols].index("canonical_url") + 1)
            if link.value:
                link.hyperlink = str(link.value)
                link.font = Font(color="1E5BC6", underline="single")
        widths = {"title": 60, "organization_name": 30, "canonical_url": 45, "notes": 50}
        for i, (k, _label) in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 16)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    @staticmethod
    def _pdf(rows: list[dict[str, Any]], scope: str, lang: str) -> bytes:
        from xml.sax.saxutils import escape

        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        fr = lang != "en"
        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=landscape(A4),
            title="Opportunex",
            author="Opportunex",
            leftMargin=1.5 * cm,
            rightMargin=1.5 * cm,
            topMargin=1.5 * cm,
            bottomMargin=1.5 * cm,
        )
        st = getSampleStyleSheet()
        story: list[Any] = [
            Spacer(1, 4 * cm),
            Paragraph("Opportunex", st["Title"]),
            Paragraph(("Rapport d'opportunités" if fr else "Opportunity report") + f" — {scope}", st["Heading2"]),
            Paragraph(date.today().isoformat(), st["Normal"]),
            PageBreak(),
        ]
        upcoming = sorted([r for r in rows if r["deadline_on"]], key=lambda r: r["deadline_on"])[:5]
        story += [
            Paragraph("Synthèse" if fr else "Summary", st["Heading1"]),
            Paragraph(
                (f"{len(rows)} opportunité(s). " if fr else f"{len(rows)} opportunity(ies). ")
                + ("Prochaines échéances : " if fr else "Next deadlines: ")
                + escape(", ".join(f"{r['title'][:40]} ({r['deadline_on']})" for r in upcoming) or "—"),
                st["Normal"],
            ),
            Spacer(1, 0.5 * cm),
        ]
        small = st["BodyText"].clone("small", fontSize=8, leading=10)
        data = [
            [
                "Titre" if fr else "Title",
                "Organisme" if fr else "Organisation",
                "Date limite" if fr else "Deadline",
                "Pays" if fr else "Country",
                "Score",
            ]
        ]
        for r in rows:
            data.append(
                [
                    Paragraph(escape(r["title"][:120]), small),
                    Paragraph(escape(r["organization_name"][:60]), small),
                    r["deadline_on"],
                    r["country_code"],
                    str(r["score"]),
                ]
            )
        table = Table(data, colWidths=[11 * cm, 6 * cm, 3 * cm, 2 * cm, 2 * cm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E5BC6")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DCE2EB")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story += [table, PageBreak(), Paragraph("Fiches" if fr else "Details", st["Heading1"])]
        for r in rows:
            story += [
                Paragraph(escape(r["title"]), st["Heading3"]),
                Paragraph(escape(f"{r['organization_name']} · {r['deadline_on'] or '—'} · {r['amount'] or ''}"), small),
                Paragraph(escape(r["summary"][:800]), st["BodyText"]),
            ]
            if r["eligibility_raw"]:
                story.append(
                    Paragraph(escape(("Éligibilité : " if fr else "Eligibility: ") + r["eligibility_raw"][:800]), small)
                )
            story += [
                Paragraph(
                    f'<link href="{escape(r["canonical_url"])}">{escape(r["canonical_url"][:120])}</link>', small
                ),
                Spacer(1, 0.3 * cm),
            ]
        doc.build(story)
        return buf.getvalue()

    # ------------------------------------------------- personal data archive
    def personal_archive(self) -> tuple[str, bytes]:
        """EF-16.07: every personal record (JSON) + decrypted attachments, as a ZIP."""
        T = self.T
        data: dict[str, Any] = {
            "schema_version": JSON_SCHEMA_VERSION,
            "exported_at": now_iso(),
            "generator": f"Opportunex {ENGINE_VERSION}",
            "tables": {},
        }
        attachments: list[tuple[str, str]] = []
        with self.active.db.read() as conn:
            for name in PERSONAL_TABLES:
                tbl = T.metadata.tables[name]
                rows = [dict(r) for r in conn.execute(select(tbl)).mappings()]
                hidden = SECRET_COLUMNS.get(name, set())
                data["tables"][name] = [
                    {k: (v.hex() if isinstance(v, bytes) else v) for k, v in r.items() if k not in hidden} for r in rows
                ]
            attachments = [
                (r[0], r[1])
                for r in conn.execute(select(T.application_documents.c.file_name, T.application_documents.c.vault_path))
            ]
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("opportunex-data.json", json.dumps(data, ensure_ascii=False, indent=2, default=str))
            for i, (fname, rel) in enumerate(attachments):
                try:
                    zf.writestr(f"documents/{i:03d}-{fname}", self.active.vault.get(rel))
                except Exception:  # noqa: BLE001 - a missing vault file is listed, not fatal
                    zf.writestr(f"documents/{i:03d}-{fname}.missing.txt", "file missing from the vault")
            zf.writestr(
                "LISEZMOI.txt",
                "Export complet de vos données Opportunex (format JSON documenté, "
                "schéma 1.0). Les pièces jointes sont déchiffrées : conservez ce fichier en lieu sûr.\n",
            )
        self._record("full_data_archive", "all_personal_data", {}, len(data["tables"]), True)
        return f"opportunex-mes-donnees-{date.today().isoformat()}.zip", buf.getvalue()

    def import_json(self, payload: dict[str, Any]) -> dict[str, int]:
        """EF-16.08: restore saved opportunities (and their tracking) from an Opportunex JSON export."""
        if payload.get("schema") != "https://opportunex.app/schemas/export":
            raise ValidationFailed("not an Opportunex export", code="import_format_invalid")
        from collection.dedup import normalize_url
        from services.tracking_service import TrackingService

        T = self.T
        created = saved = 0
        tracking = TrackingService(self.active)
        for item in payload.get("items", []):
            url = str(item.get("canonical_url") or "")
            title = str(item.get("title") or "").strip()
            if not url.startswith(("http://", "https://")) or not title:
                continue
            with self.active.db.tx() as conn:
                opp_id = conn.execute(
                    select(T.opportunities.c.id).where(T.opportunities.c.normalized_url == normalize_url(url))
                ).scalar()
                if opp_id is None:
                    from collection.dedup import fingerprint

                    opp_id = uuid7()
                    conn.execute(
                        insert(T.opportunities).values(
                            id=opp_id,
                            canonical_url=url,
                            normalized_url=normalize_url(url),
                            dedup_fingerprint=fingerprint(title, item.get("organization_name"), None),
                            title=title[:300],
                            organization_name=item.get("organization_name") or None,
                            summary=item.get("summary") or None,
                            deadline_on=item.get("deadline_on") or None,
                            eligibility_raw=item.get("eligibility_raw") or None,
                            verification_status="unverified",
                        )
                    )
                    created += 1
                already = conn.execute(
                    select(T.saved_opportunities.c.id).where(
                        T.saved_opportunities.c.user_id == self.uid, T.saved_opportunities.c.opportunity_id == opp_id
                    )
                ).first()
            if already is None and ("status" in item or payload.get("scope") in ("saved", "applications")):
                tracking.save(opp_id, item.get("priority") or "medium")
                saved += 1
        return {"opportunities_created": created, "saved": saved}
