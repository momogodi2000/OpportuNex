from __future__ import annotations

import io
import json
import time
import zipfile
from datetime import UTC, datetime, timedelta, timezone

import pytest

from services.monitor_service import MonitorService, next_run, validate_schedule
from tests.conftest import EngineClient, add_source, collect
from tests.fixtures.site import FixtureSite


def test_dashboard_empty_then_populated(client: EngineClient, site: FixtureSite) -> None:
    empty = client.get("/dashboard").json()
    assert empty["is_new_user"] is True and empty["counters"]["saved"] == 0
    items = collect(client, site)
    master = next(i for i in items if i["title"].startswith("Bourse"))
    client.post("/saved", json={"opportunity_id": master["id"]})
    d = client.get("/dashboard").json()
    assert d["counters"]["saved"] == 1 and d["counters"]["discovered"] >= 3
    assert d["counters"]["new_since_last_search"] == 3
    assert d["applications"]["by_status"] == {"discovered": 1}
    assert d["activity"][0]["results_count"] == 3
    assert d["sources"]["total"] >= 70
    assert all(r["id"] != master["id"] for r in d["recommendations"])
    assert d["is_new_user"] is False


def test_schedule_rules() -> None:
    base = datetime(2026, 10, 15, 9, 0, tzinfo=UTC)
    assert next_run("daily", "08:00", base, UTC) == datetime(2026, 10, 16, 8, 0, tzinfo=UTC)
    assert next_run("weekly", "mon 07:30", base, UTC).weekday() == 0
    assert next_run("custom_cron", "every 48h", base) == datetime(2026, 10, 17, 9, 0, tzinfo=UTC)
    # Regression: "08:00" is the user's local time, not UTC (Douala is UTC+1).
    douala = timezone(timedelta(hours=1))
    assert next_run("daily", "08:00", base, douala) == datetime(2026, 10, 16, 7, 0, tzinfo=UTC)
    assert next_run("daily", "11:00", base, douala) == datetime(2026, 10, 15, 10, 0, tzinfo=UTC)
    monday = next_run("weekly", "lun 07:30", base, douala)
    assert monday.astimezone(douala).weekday() == 0 and monday.astimezone(douala).hour == 7
    from core.errors import ValidationFailed

    with pytest.raises(ValidationFailed):
        validate_schedule("custom_cron", "every 2h")
    with pytest.raises(ValidationFailed):
        validate_schedule("daily", "25:99")


def test_monitor_run_notifies_new_relevant_items(client: EngineClient, site: FixtureSite) -> None:
    source_id = add_source(client, site.url("/feed.xml"), "Fixture")
    m = client.post(
        "/monitors",
        json={
            "name": "Bourses",
            "filters": {"keywords": []},
            "schedule_kind": "daily",
            "schedule_expr": "08:00",
            "min_score": 0,
            "source_ids": [source_id],
        },
    )
    assert m.status_code == 201, m.text
    mid = m.json()["id"]
    assert (
        client.post(
            "/monitors", json={"name": "x", "filters": {}, "schedule_kind": "custom_cron", "schedule_expr": "every 1h"}
        ).json()["code"]
        == "monitor_too_frequent"
    )
    assert client.post(f"/monitors/{mid}/run").status_code == 202
    for _ in range(150):
        runs = client.get(f"/monitors/{mid}").json()["runs"]
        if runs and runs[0]["status"] != "running":
            break
        time.sleep(0.2)
    assert runs[0]["status"] == "completed" and runs[0]["new_items"] == 3 and runs[0]["notified"] == 1
    notes = client.get("/notifications", params={"unread": True}).json()["items"]
    assert notes[0]["kind"] == "new_opportunities"
    client.post("/notifications/read", json={})
    assert client.get("/notifications", params={"unread": True}).json()["items"] == []
    client.patch(f"/monitors/{mid}", json={"enabled": False})
    assert client.get(f"/monitors/{mid}").json()["enabled"] is False
    assert client.get("/monitors/history").json()["items"][0]["name"] == "Bourses"
    assert client.delete(f"/monitors/{mid}").status_code == 204


def test_due_reminders_become_notifications(client: EngineClient, site: FixtureSite) -> None:
    master = next(i for i in collect(client, site) if i["title"].startswith("Bourse"))
    s = client.post("/saved", json={"opportunity_id": master["id"]}).json()
    active = client.state.active
    T = active.tables
    from sqlalchemy import update

    with active.db.tx() as conn:
        conn.execute(
            update(T.reminders)
            .where(T.reminders.c.saved_id == s["id"], T.reminders.c.offset_days == 14)
            .values(remind_at="2000-01-01T00:00:00Z")
        )
    client.put("/notifications/preferences/deadline", json={"enabled": True})
    sent = MonitorService(active, client.state).due_reminders()
    assert sent == 1
    assert client.get("/notifications").json()["items"][0]["body"] == "J-14"
    client.put("/notifications/reminder-offsets", json={"offsets": [10, 2]})
    assert client.get("/notifications/preferences").json()["items"][0]["reminder_offsets"] == [10, 2]


def test_workspaces_templates_and_compare(client: EngineClient, site: FixtureSite) -> None:
    items = collect(client, site)
    templates = client.get("/workspaces/templates").json()["items"]
    assert len(templates) == 7
    ws = client.post("/workspaces", json={"name": "Bourses 2027", "template_code": "master_scholarship"}).json()
    assert ws["suggested_filters"]["categories"] == ["scholarship.master"]
    assert client.post("/workspaces", json={"name": "Bourses 2027"}).status_code == 409
    ws = client.put(f"/workspaces/{ws['id']}/opportunities/{items[0]['id']}").json()
    assert len(ws["opportunities"]) == 1
    ws = client.post(f"/workspaces/{ws['id']}/notes", json={"body": "Priorité Europe"}).json()
    assert ws["notes"][0]["body"] == "Priorité Europe"
    assert client.patch(f"/workspaces/{ws['id']}", json={"archived": True}).json()["archived"] is True
    assert client.get("/workspaces").json()["items"] == []
    assert len(client.get("/workspaces", params={"archived": True}).json()["items"]) == 1
    cmp = client.post("/opportunities/compare", json={"ids": [items[0]["id"], items[1]["id"]]}).json()["items"]
    assert len(cmp) == 2 and "education_required" in cmp[0]
    assert client.post("/opportunities/compare", json={"ids": [items[0]["id"]]}).status_code == 422


def test_exports_all_formats(client: EngineClient, site: FixtureSite) -> None:
    items = collect(client, site)
    run_id = client.get("/search/history").json()["items"][0]["last_run"]["id"]
    csv = client.post("/exports", json={"format": "csv", "scope": "filtered_results", "run_id": run_id})
    assert csv.status_code == 200 and csv.content.startswith("﻿".encode())
    assert "Titre;Organisme" in csv.content.decode("utf-8-sig")
    assert "Notes" not in csv.content.decode("utf-8-sig")  # no personal data by default (EF-16.06)
    xlsx = client.post("/exports", json={"format": "xlsx", "scope": "selection", "ids": [items[0]["id"]]})
    assert xlsx.content.startswith(b"PK")
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws.auto_filter.ref and ws.max_row == 2
    pdf = client.post("/exports", json={"format": "pdf", "scope": "filtered_results", "run_id": run_id})
    assert pdf.content.startswith(b"%PDF")
    js = client.post("/exports", json={"format": "json", "scope": "filtered_results", "run_id": run_id}).json()
    assert js["schema_version"] == "1.0" and len(js["items"]) == 3
    assert "attachment;" in csv.headers["content-disposition"]


def test_personal_archive_and_json_import(client: EngineClient, site: FixtureSite) -> None:
    items = collect(client, site)
    s = client.post("/saved", json={"opportunity_id": items[0]["id"]}).json()
    client.post(f"/saved/{s['id']}/documents", files={"file": ("cv.txt", b"mon cv")}, data={"label": "CV"})
    z = client.post("/exports/personal-data")
    with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
        data = json.loads(zf.read("opportunex-data.json"))
        assert "password_hash" not in data["tables"]["users"][0]
        assert data["tables"]["saved_opportunities"][0]["id"] == s["id"]
        doc = next(n for n in zf.namelist() if n.startswith("documents/"))
        assert zf.read(doc) == b"mon cv"
    exported = client.post("/exports", json={"format": "json", "scope": "saved"}).content
    client.delete(f"/saved/{s['id']}")
    r = client.post("/imports/json", files={"file": ("x.json", exported)})
    assert r.json()["saved"] == 1
    assert len(client.get("/saved").json()["items"]) == 1
    bad = client.post("/imports/json", files={"file": ("x.json", b"{}")})
    assert bad.status_code == 422


def test_privacy_dashboard_consents_and_erase(client: EngineClient, site: FixtureSite) -> None:
    collect(client, site)
    p = client.get("/privacy").json()
    assert p["telemetry"] is False and p["stored"]["opportunities"] == 3
    terms = next(c for c in p["consents"] if c["purpose"] == "terms")
    assert client.post(f"/privacy/consents/{terms['id']}/revoke").json()["code"] == "consent_required"
    assert client.post("/privacy/erase", json={"kinds": ["search_history"], "confirmation": "NO"}).status_code == 422
    r = client.post(
        "/privacy/erase",
        json={"kinds": ["search_history", "unsaved_opportunities", "profile_details"], "confirmation": "ERASE"},
    ).json()
    assert r["unsaved_opportunities"] == 3
    after = client.get("/privacy").json()["stored"]
    assert after["opportunities"] == 0 and after["search_history"] == 0 and after["profile"]["skills"] == 0


def test_auto_archive_delay_is_configurable(client: EngineClient, site: FixtureSite) -> None:
    from sqlalchemy import update

    from services.tracking_service import TrackingService

    opp = collect(client, site)[0]
    saved = client.post("/saved", json={"opportunity_id": opp["id"]}).json()
    app_id = saved["application"]["id"]
    client.post(f"/applications/{app_id}/status", json={"status": "withdrawn"})
    active = client.state.active
    with active.db.tx() as conn:
        conn.execute(
            update(active.tables.applications)
            .where(active.tables.applications.c.id == app_id)
            .values(updated_at="2000-01-01T00:00:00Z")
        )
    assert client.get("/settings/tracking").json() == {"auto_archive_days": 60}
    assert client.put("/settings/tracking", json={"auto_archive_days": -1}).status_code == 422
    client.put("/settings/tracking", json={"auto_archive_days": 0})
    assert TrackingService(active).auto_archive() == 0  # 0 turns the sweep off
    client.put("/settings/tracking", json={"auto_archive_days": 30})
    assert TrackingService(active).auto_archive() == 1
    assert client.get(f"/saved/{saved['id']}").json()["archived"] is True


def test_network_settings(client: EngineClient) -> None:
    assert client.get("/settings/network").json() == {"proxy_url": None, "data_saver": False, "max_mb_per_search": 50}
    assert client.put("/settings/network", json={"proxy_url": "ftp://x"}).status_code == 422
    r = client.put(
        "/settings/network", json={"proxy_url": "http://proxy.local:3128", "data_saver": True, "max_mb_per_search": 20}
    ).json()
    assert r == {"proxy_url": "http://proxy.local:3128", "data_saver": True, "max_mb_per_search": 20}
